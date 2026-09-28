#!/usr/bin/env python3
"""ARM450_KINEMATICS_slides.pptx -- the presentation for the guide (as the
myCobot thesis deck): one idea per slide, the analysis figures, the equation
pages of the report rendered as images (their prescript notation survives),
the sine animation (GIFs animate in slideshow)."""
import os
import re

import pymupdf as fitz
from pptx import Presentation
from pptx.util import Inches, Pt

DOC = os.path.dirname(os.path.abspath(__file__))
KIN = os.path.dirname(DOC)
FIG = os.path.join(KIN, "figures")
OUT = os.path.join(DOC, "ARM450_KINEMATICS_slides.pptx")


def txt(fn):
    p = os.path.join(KIN, fn)
    return open(p).read() if os.path.exists(p) else ""


def report_page_png(i, name):
    d = fitz.open(os.path.join(DOC, "ARM450_KINEMATICS.pdf"))
    p = os.path.join(FIG, "_slide_%s.png" % name)
    d[i].get_pixmap(dpi=200).save(p)
    return p


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def slide(title, img=None, bullets=None, notes="", img_box=(0.3, 1.05, 12.7, 6.2), bfs=18):
        s = prs.slides.add_slide(blank)
        tb = s.shapes.add_textbox(Inches(0.4), Inches(0.2), Inches(12.5), Inches(0.8))
        p = tb.text_frame.paragraphs[0]; p.text = title; p.font.size = Pt(30); p.font.bold = True
        if bullets:
            bx = s.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(12.3), Inches(5.9))
            tf = bx.text_frame; tf.word_wrap = True
            for k, b in enumerate(bullets):
                q = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                q.text = b; q.font.size = Pt(bfs); q.space_after = Pt(8)
        if img:
            x, y, w, h = img_box
            from PIL import Image
            iw, ih = Image.open(img).size
            sc = min(w / iw, h / ih)
            s.shapes.add_picture(img, Inches(x + (w - iw * sc) / 2), Inches(y + (h - ih * sc) / 2),
                                 Inches(iw * sc), Inches(ih * sc))
        if notes:
            s.notes_slide.notes_text_frame.text = notes
        return s

    fk, ws, cmp_, sine = txt("VERIFY_FK.log"), txt("WORKSPACE.log"), txt("IK_COMPARISON.log"), txt("SINE_DEMO.log")
    g = lambda rx, s: (re.search(rx, s).group(1) if re.search(rx, s) else "?")
    slide("ARM-450 rev I.1 -- kinematics", bullets=[
        "DH table, forward and inverse kinematics, Jacobian, workspace -- in my convention:",
        "Modified (Craig) DH + velocity-propagation Jacobian (after the myCobot 280 thesis and Rsine)",
        "Geometry from the rev I.1 CAD: d1 = 90, a2 = 119, d4 = 221.369, d_T = 88.089 mm; spherical wrist",
        "FK = CAD to %s mm, = URDF to %s mm (2000 poses)" % (g(r"product of exponentials\), position\s+max ([\d.e+-]+)", fk),
                                                            g(r"URDF \(ikpy [\d.]+\), position\s+max ([\d.e+-]+)", fk)),
        "Closed-form IK: 8 branches, exact, 0.2 ms; numerical solvers compared",
        "Pen straight down impossible (172.5 deg of pitch) -> tilted pen on the table; vertical board works"],
        notes="Motivation: same method as the myCobot thesis, applied to the arm I designed and printed.")
    slide("My convention: Modified (Craig) DH", img=report_page_png(1, "conv"),
          notes="Frame i on link i; a and alpha measured along/about X_{i-1}; d and theta along/about Z_i.")
    slide("DH table (derived from the CAD servo axes)", img=os.path.join(FIG, "fig_dh_table.png"),
          notes="Every link transform read back from frames built by the rules; exact Craig form; matches the model to 1e-14 mm.")
    slide("Frames and link lengths", img=os.path.join(FIG, "fig_dh_frames.png"),
          notes="d4 = 221.369 is elbow to wrist centre: the wrist centre is where J4, J5, J6 meet.")
    slide("Forward kinematics -- verified three ways", img=os.path.join(FIG, "fig_fk_verification.png"),
          notes="CAD by product of exponentials of the measured servo axes; URDF through ikpy; Jacobian vs z x r and finite differences.")
    slide("Jacobian by velocity propagation", img=report_page_png(6, "jac"))
    slide("Inverse kinematics in closed form", img=report_page_png(7, "ik"),
          notes="Wrist centre decouples position and orientation. 2 x 2 x 2 branches, then the joint limits.")
    slide("The 8 branches of one pose", img=os.path.join(FIG, "fig_ik_branches.png"))
    slide("IK solvers compared (thesis Exp. 2)", img=os.path.join(FIG, "fig_ik_comparison.png"),
          notes=cmp_)
    slide("Reachable workspace", img=os.path.join(FIG, "fig_workspace.png"),
          notes=ws.split("TASK-CONSTRAINED")[0])
    slide("Task-constrained workspace (thesis Exp. 1)", img=os.path.join(FIG, "fig_task_workspace.png"),
          notes="Straight down needs 180 deg of pitch; J2 + J3 + J5 give 172.5. Table: tilted pen on a ring. Board: front pen only.")
    slide("Manipulability and singularities (thesis Exp. 3)", img=os.path.join(FIG, "fig_manipulability.png"),
          notes="Home is singular three ways: shoulder, elbow, wrist. Work from the ready pose.")
    slide("Sine tracing (after Rsine)", img=os.path.join(FIG, "fig_sine_demos.png"), notes=sine)
    tim = txt("TIMING.log")
    slide("The sine on the real servos: timing and encoder floor", img=os.path.join(FIG, "fig_timing.png"), notes=tim)
    s = slide("The CAD arm drawing", img=os.path.join(FIG, "fig_sine_cad.png"), img_box=(0.3, 1.0, 12.7, 3.3))
    for k, gif in enumerate(("sine_vertical.gif", "sine_table.gif")):
        p = os.path.join(FIG, gif)
        if os.path.exists(p):
            s.shapes.add_picture(p, Inches(2.2 + 5.2 * k), Inches(4.3), Inches(3.1), Inches(3.1))
    slide("Conclusions", bullets=[
        "DH table follows from the CAD with my frame rules; FK = CAD = URDF to 1e-13 mm",
        "Velocity-propagation Jacobian = geometric form; manipulability / conditioning mapped",
        "Spherical wrist -> closed-form IK (exact, all branches); numerical solvers are local: use them along paths",
        "Home is triply singular: start from the bent ready pose",
        "Joint limits cap the pitch at 172.5 deg: no pen straight down; tilted pen on a ring; board from ~190 mm",
        "Free pen roll keeps the traced motion smooth (largest step < 3 deg) -- redundancy, as in the thesis",
        "On the ST3215: board sine at ~80 mm/s with the servos at half their rating; encoder floor ~0.9 mm"])
    prs.save(OUT)
    for f in os.listdir(FIG):
        if f.startswith("_slide_"):
            os.remove(os.path.join(FIG, f))
    print("written", OUT, len(prs.slides), "slides")


if __name__ == "__main__":
    main()
