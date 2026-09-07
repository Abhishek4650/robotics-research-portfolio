#!/usr/bin/env python3
"""
Detailed thesis slide deck, with the locked 'paper' drawing animation embedded on
the final slide (animates during a PowerPoint/Impress slideshow).

Run:  python3 analysis/make_detailed_ppt.py -> docs/mycobot_thesis_detailed.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

PKG = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(PKG, "figures")
OUT = os.path.join(PKG, "docs", "mycobot_thesis_detailed.pptx")
NAVY, ACCENT, GREY = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0xD6, 0x27, 0x28), RGBColor(0x33, 0x33, 0x33)

S = [
    dict(title="Task-Constrained Workspace & Inverse Kinematics\nfor Surface-Tracing with the myCobot 280",
         sub="A detailed study — where, how, and with what posture a 6-DOF arm can draw",
         bullets=[], fig=None),
    dict(title="Motivation", sub="Two questions from the drawing project",
         bullets=["Reaching a point ≠ reaching it with the pen perpendicular to the surface",
                  "Q1: which pen directions are feasible where? (front vs behind)",
                  "Q2: can it trace on any surface anywhere — incl. a horizontal table?",
                  "Thesis: the reachable and the DRAWABLE workspaces are different sets"],
         fig=None),
    dict(title="Contributions", sub=None,
         bullets=["A verified modified-DH model + velocity-propagation Jacobian",
                  "A map of the task-constrained (drawable) workspace, vertical & horizontal",
                  "A comparison of four inverse-kinematics solvers",
                  "Manipulability / conditioning maps explaining the drawable region",
                  "Redundancy resolution: same pen path, chosen arm posture",
                  "Live ROS 2 / RViz demonstrations"],
         fig=None),
    dict(title="Kinematic Model — Modified (Craig) DH", sub="Verified to ~1e-15 m vs the URDF",
         bullets=["Link transform:  Aᵢ = Rotx(αᵢ₋₁)·Transx(aᵢ₋₁)·Rotz(θᵢ)·Transz(dᵢ)",
                  "Base-to-tool:  T = A₁A₂A₃A₄A₅A₆",
                  "DH table identified by fitting to the URDF (machine-precision match)",
                  "αᵢ₋₁ link twist · aᵢ₋₁ link length · dᵢ offset · θᵢ joint angle"],
         fig=None),
    dict(title="Jacobian & Dexterity Metrics", sub="Velocity-propagation method",
         bullets=["Column i:  Jᵢ = [ zᵢ × (pₑ − pᵢ) ; zᵢ ]   (verified vs finite diff 2.5e-7)",
                  "Twist:  [v; ω] = J(q)·q̇",
                  "Manipulability:  w = √det(Jᵥ Jᵥᵀ)   (0 at a singularity)",
                  "Condition number:  κ = σ_max / σ_min   (1 = isotropic)"],
         fig=None),
    dict(title="Task-Constrained Workspace — Definition", sub=None,
         bullets=["Reachable set R = positions attainable within joint limits",
                  "Drawable set D(n) = positions reachable with pen axis ẑ = ±n (surface normal)",
                  "Estimated by testing DLS IK convergence with the perpendicular-pen constraint",
                  "Key idea: D(n) ⊂ R, and the gap depends on the surface orientation"],
         fig=None),
    dict(title="Feasibility Results", sub="Where the pen can be held normal to the surface",
         bullets=["Vertical: 57.9% of R allows a ⊥ pen; only 35.2% front-facing",
                  "Near base → only BACK; outer region → FRONT opens up",
                  "Horizontal table: 63.5% drawable pen-down — an ANNULUS",
                  "Dead zone directly above the base (wrist can't point straight down)"],
         fig="fig_feasibility.png"),
    dict(title="Inverse-Kinematics Comparison", sub="Four Jacobian-based solvers",
         bullets=["Damped least squares: 98.7% success — most robust",
                  "Pseudoinverse / Levenberg–Marquardt: quadratic but brittle near singularities",
                  "Jacobian-transpose: always progresses, ~25× slower",
                  "→ DLS is the right default for surface tracing"],
         fig="fig_ik_comparison.png"),
    dict(title="Manipulability & Conditioning", sub="Why the drawable region has its shape",
         bullets=["Low-manipulability dead zone directly above the base",
                  "Bright dexterous RING at mid-radius — best drawing region",
                  "Condition number rises toward the workspace boundary",
                  "The drawable annulus = the high-w, low-κ ring"],
         fig="fig_manipulability.png"),
    dict(title="Demonstrations", sub="Sine tracing on two surfaces (0.000 mm)",
         bullets=["Placements chosen from the feasibility & manipulability maps",
                  "Front-facing vertical board — pen points away from the robot",
                  "Horizontal table — pen points down, in the dexterous ring"],
         fig="fig_demos.png"),
    dict(title="Redundancy — Same Path, Chosen Posture", sub="Same end-effector, arm 'opened' toward z",
         bullets=["The 6-DOF arm reaches the same pen poses with different postures (IK branches)",
                  "We select a RAISED (elbow-up) branch: upper arm 0.10 m → 0.17 m",
                  "End-effector traces the IDENTICAL sine (EE z = 0.100 m unchanged)",
                  "Smooth along the path (max 3°/frame, no posture jumps)"],
         fig="fig_posture_compare.png"),
    dict(title="Conclusions & Future Work", sub=None,
         bullets=["Reachable ≠ drawable; the gap depends on the surface normal",
                  "Vertical 57.9% / horizontal 63.5% drawable; feasibility ↔ manipulability ring",
                  "DLS is the most robust IK; redundancy sets the posture",
                  "Next: tilted surfaces, manipulability-optimal drawing, real robot + Gazebo"],
         fig="fig_sine_3d.png" if os.path.exists(os.path.join(FIG, "fig_sine_3d.png")) else None),
    dict(title="Live Demonstration — Writing Left-to-Right", sub="ROS 2 / RViz (animation plays in slideshow)",
         bullets=["The locked 'paper' demo: front table, raised posture, left-to-right sine",
                  "Command:  ros2 launch mycobot_thesis thesis_draw.launch.py surface:=paper"],
         fig="paper_raised_draw.gif", animation=True),
]


def _title(slide, text, size=26):
    b = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(1.1))
    p = b.text_frame.paragraphs[0]; b.text_frame.word_wrap = True
    p.text = text; p.font.size = Pt(size); p.font.bold = True; p.font.color.rgb = NAVY


def build():
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for i, s in enumerate(S):
        slide = prs.slides.add_slide(blank)
        img = os.path.join(FIG, s["fig"]) if s.get("fig") else None
        if i == 0:
            t = slide.shapes.add_textbox(Inches(0.8), Inches(2.4), Inches(11.7), Inches(2.0))
            t.text_frame.word_wrap = True
            p = t.text_frame.paragraphs[0]; p.text = s["title"]
            p.font.size = Pt(30); p.font.bold = True; p.font.color.rgb = NAVY; p.alignment = PP_ALIGN.CENTER
            sb = slide.shapes.add_textbox(Inches(0.8), Inches(4.6), Inches(11.7), Inches(0.8))
            sp = sb.text_frame.paragraphs[0]; sp.text = s["sub"]; sp.font.size = Pt(17)
            sp.font.color.rgb = ACCENT; sp.alignment = PP_ALIGN.CENTER
            continue
        _title(slide, s["title"])
        if s.get("sub"):
            sb = slide.shapes.add_textbox(Inches(0.5), Inches(1.15), Inches(12.3), Inches(0.5))
            sp = sb.text_frame.paragraphs[0]; sp.text = s["sub"]; sp.font.size = Pt(15)
            sp.font.italic = True; sp.font.color.rgb = ACCENT
        if s.get("animation") and img and os.path.exists(img):
            bx = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(5.6), Inches(5.0))
            for k, b in enumerate(s["bullets"]):
                p = bx.text_frame.paragraphs[0] if k == 0 else bx.text_frame.add_paragraph()
                p.text = "•  " + b; p.font.size = Pt(15); p.font.color.rgb = GREY; p.space_after = Pt(8)
            bx.text_frame.word_wrap = True
            # A GIF added as a picture animates during a PowerPoint / Impress slideshow.
            slide.shapes.add_picture(img, Inches(6.4), Inches(1.9), width=Inches(6.4))
        else:
            bx = slide.shapes.add_textbox(Inches(0.5), Inches(1.8),
                                          Inches(6.0 if img else 12.2), Inches(5.2))
            for k, b in enumerate(s["bullets"]):
                p = bx.text_frame.paragraphs[0] if k == 0 else bx.text_frame.add_paragraph()
                p.text = "•  " + b; p.font.size = Pt(15); p.font.color.rgb = GREY; p.space_after = Pt(8)
            bx.text_frame.word_wrap = True
            if img:
                slide.shapes.add_picture(img, Inches(6.7), Inches(1.9), width=Inches(6.3))
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    print("wrote", OUT, f"({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    build()
