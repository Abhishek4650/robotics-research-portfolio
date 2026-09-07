#!/usr/bin/env python3
"""
Thesis slide deck (PowerPoint) from the experiment figures + findings.

Run:  python3 analysis/make_thesis_ppt.py -> docs/mycobot_thesis_slides.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

PKG = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIG = os.path.join(PKG, "figures")
OUT = os.path.join(PKG, "docs", "mycobot_thesis_slides.pptx")
NAVY, ACCENT, GREY = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0xD6, 0x27, 0x28), RGBColor(0x33, 0x33, 0x33)

SLIDES = [
    dict(title="Task-Constrained Workspace & Inverse Kinematics\nfor Surface-Tracing with the myCobot 280",
         subtitle="Where — and how well — can a 6-DOF arm draw on a surface?",
         bullets=[], figure=None, note=(
             "Open with the core idea: reaching a point is not enough — the tool must "
             "be held normal to the surface, and that shrinks the usable workspace.")),
    dict(title="Motivation", subtitle="Two questions from the drawing project",
         bullets=[
             "Reaching a point ≠ reaching it with the pen perpendicular to the surface",
             "Q1: which pen directions are feasible where? (front vs behind)",
             "Q2: can it trace on any surface anywhere — incl. a horizontal table?",
             "Thesis: the reachable and the DRAWABLE workspaces are different sets"],
         figure=None, note=(
             "These two questions came from the myCobot drawing project and turn out to "
             "be a real research problem about task-constrained workspaces.")),
    dict(title="Kinematic Model", subtitle="Modified-DH, verified to 1e-15 m",
         bullets=[
             "Modified (Craig) DH: link transform Aᵢ = Rotx(α)·Transx(a)·Rotz(θ)·Transz(d)",
             "DH table identified from the URDF; FK matches URDF to ~1e-15 m",
             "Velocity-propagation Jacobian: Jᵢ = [ zᵢ×(pₑ−pᵢ) ; zᵢ ], verified 2.5e-7",
             "Dexterity: manipulability w=√det(Jv Jvᵀ), condition number κ"],
         figure=None, note=(
             "Emphasize the model is verified, not assumed — FK matches the manufacturer "
             "model to machine precision.")),
    dict(title="Task-Constrained Workspace — Results",
         subtitle="Where the pen can be held normal to the surface",
         bullets=[
             "Vertical: only 57.9% of reachable space allows a ⊥ pen; 35.2% front-facing",
             "Near the base only BACK works; FRONT opens up in the outer region",
             "Horizontal table: 63.5% drawable pen-down — an ANNULUS (dead zone above base)"],
         figure="fig_feasibility.png", note=(
             "Panel (a): yellow = back-only near base, red = front-only farther out — that's "
             "the answer to the front/back question. Panels (b,c): the drawable ring on a table.")),
    dict(title="Inverse-Kinematics Comparison",
         subtitle="Four Jacobian-based solvers on the tracing task",
         bullets=[
             "Damped least squares: 98.7% success — most robust (damping caps accuracy)",
             "Pseudoinverse / Levenberg–Marquardt: quadratic but brittle near singularities",
             "Jacobian-transpose: always progresses, but ~25× slower",
             "→ DLS is the right default for well-seeded surface tracing"],
         figure="fig_ik_comparison.png", note=(
             "Panel (d) shows convergence rate on a good pose; the bars show robustness. The "
             "story is accuracy vs robustness — DLS wins on robustness.")),
    dict(title="Manipulability & Conditioning",
         subtitle="Why the drawable region has its shape",
         bullets=[
             "Low-manipulability 'dead zone' directly above the base",
             "Bright dexterous RING at mid-radius — best place to draw",
             "Condition number rises toward the workspace boundary (near-singular)",
             "The drawable annulus coincides with the high-w, low-κ ring"],
         figure="fig_manipulability.png", note=(
             "This explains the annulus in the feasibility map: feasibility and good "
             "conditioning are geometrically the same region.")),
    dict(title="Demonstrations", subtitle="Sine tracing on two surfaces (0.000 mm)",
         bullets=[
             "Placements chosen using the feasibility & manipulability maps",
             "(a) FRONT-facing vertical board — pen points away from the robot",
             "(b) HORIZONTAL table — pen points down, in the dexterous ring",
             "Both track the target to 0.000 mm"],
         figure="fig_demos.png", note=(
             "The maps predict where clean drawing is possible; these demos confirm it, "
             "including the front-facing and the horizontal cases.")),
    dict(title="Conclusions & Future Work", subtitle=None,
         bullets=[
             "Reachable ≠ drawable; the gap depends on the surface normal",
             "Vertical 57.9% / horizontal 63.5% drawable; feasibility ↔ manipulability ring",
             "DLS is the most robust IK for the task",
             "Next: tilted surfaces, redundancy for manipulability-optimal drawing, real robot + Gazebo"],
         figure=None, note=(
             "Close on the contribution: a method to map the drawable workspace and place "
             "drawing tasks where they succeed.")),
]


def _title(slide, text, size=28):
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.3), Inches(1.1))
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text
    p.font.size = Pt(size); p.font.bold = True; p.font.color.rgb = NAVY


def build():
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    for i, s in enumerate(SLIDES):
        slide = prs.slides.add_slide(blank)
        img = os.path.join(FIG, s["figure"]) if s["figure"] else None
        if i == 0:
            t = slide.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.7), Inches(2.0))
            tf = t.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]; p.text = s["title"]
            p.font.size = Pt(32); p.font.bold = True; p.font.color.rgb = NAVY
            p.alignment = PP_ALIGN.CENTER
            sub = slide.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(11.7), Inches(0.8))
            sp = sub.text_frame.paragraphs[0]; sp.text = s["subtitle"]
            sp.font.size = Pt(18); sp.font.color.rgb = ACCENT; sp.alignment = PP_ALIGN.CENTER
        else:
            _title(slide, s["title"])
            if s.get("subtitle"):
                sb = slide.shapes.add_textbox(Inches(0.5), Inches(1.15), Inches(12.3), Inches(0.5))
                sp = sb.text_frame.paragraphs[0]; sp.text = s["subtitle"]
                sp.font.size = Pt(15); sp.font.italic = True; sp.font.color.rgb = ACCENT
            bx = slide.shapes.add_textbox(Inches(0.5), Inches(1.8),
                                          Inches(6.0 if img else 12.2), Inches(5.2))
            tf = bx.text_frame; tf.word_wrap = True
            for k, b in enumerate(s["bullets"]):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                p.text = "•  " + b; p.font.size = Pt(16); p.font.color.rgb = GREY
                p.space_after = Pt(8)
            if img:
                slide.shapes.add_picture(img, Inches(6.7), Inches(1.9), width=Inches(6.3))
        slide.notes_slide.notes_text_frame.text = s["note"]
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    print("wrote", OUT, f"({len(prs.slides._sldIdLst)} slides)")


if __name__ == "__main__":
    build()
